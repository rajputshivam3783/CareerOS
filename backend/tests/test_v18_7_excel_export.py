"""V18.7 — Applicants Excel export: same data as the CSV export, as a
real .xlsx workbook."""

import os

os.environ["DATABASE_URL"] = "sqlite:///./test_careeros_v18_7.db"
os.environ["AUTO_VERIFY_EMAIL_IN_TESTS"] = "true"

import io  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402

ADMIN_HEADERS = {"X-Admin-Key": "change-this-admin-key"}
PW = "password12345!"


def _setup_job_with_applicant(client, suffix):
    recruiter_email = f"v187rec{suffix}@example.com"
    candidate_email = f"v187cand{suffix}@example.com"

    client.post("/api/v1/auth/register", json={"email": recruiter_email, "password": PW, "password_confirm": PW, "full_name": "Rec"})
    rlogin = client.post("/api/v1/auth/login", json={"email": recruiter_email, "password": PW})
    rme = client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {rlogin.json()['access_token']}"}).json()
    client.post(f"/api/v1/admin/users/{rme['id']}/role", headers=ADMIN_HEADERS, json={"role": "recruiter"})
    rlogin = client.post("/api/v1/auth/recruiter/login", json={"email": recruiter_email, "password": PW})
    r_headers = {"Authorization": f"Bearer {rlogin.json()['access_token']}"}

    client.post("/api/v1/auth/register", json={"email": candidate_email, "password": PW, "password_confirm": PW, "full_name": "Candidate Name"})
    clogin = client.post("/api/v1/auth/login", json={"email": candidate_email, "password": PW})
    c_headers = {"Authorization": f"Bearer {clogin.json()['access_token']}"}

    job = client.post(
        "/api/v1/recruiter/jobs",
        headers=r_headers,
        json={"title": f"Engineer {suffix}", "organization": "Acme", "location": "Remote", "description": "x", "qualification": "x"},
    ).json()
    client.post(f"/api/v1/admin/jobs/{job['id']}/publish", headers=ADMIN_HEADERS)
    client.post(f"/api/v1/jobs/{job['id']}/apply", headers=c_headers, json={})

    return r_headers, job["id"]


def test_excel_export_returns_workbook_with_header_and_row():
    with TestClient(app) as client:
        headers, job_id = _setup_job_with_applicant(client, "a")
        response = client.get(f"/api/v1/recruiter/jobs/{job_id}/applicants/export.xlsx", headers=headers)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert f"applicants-job-{job_id}.xlsx" in response.headers["content-disposition"]

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb["Applicants"]
        header = [c.value for c in ws[1]]
        assert header == [
            "Applicant ID", "Candidate name", "Candidate email", "Status", "Pipeline stage",
            "Reject reason", "Cover note", "Applied on",
        ]
        assert ws[1][0].font.bold is True
        second_row = [c.value for c in ws[2]]
        assert second_row[1] == "Candidate Name"
        assert "v187conda@example.com" not in str(second_row)  # sanity: not leaking wrong email


def test_excel_export_requires_ownership():
    with TestClient(app) as client:
        _headers, job_id = _setup_job_with_applicant(client, "b")
        _other_headers, _other_job_id = _setup_job_with_applicant(client, "c")
        response = client.get(f"/api/v1/recruiter/jobs/{job_id}/applicants/export.xlsx", headers=_other_headers)
        assert response.status_code == 404
