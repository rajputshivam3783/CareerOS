"""V8 — resume text extraction and skill detection.

Supports PDF and DOCX (the two formats candidates overwhelmingly use).
Extraction is best-effort, same honesty stance as the V5 notification
PDF parser: a scanned-image resume with no text layer returns an empty
string rather than a fabricated guess, and the caller surfaces that
plainly instead of pretending it read something it didn't.
"""

from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.services.career import SKILLS

_MAX_TEXT_CHARS = 20_000  # generous for a resume; keeps storage/matching bounded


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    name = (filename or "").lower()

    if name.endswith(".pdf"):
        text = _extract_pdf(file_bytes)
    elif name.endswith(".docx"):
        text = _extract_docx(file_bytes)
    else:
        # Best-effort: try plain-text decoding (e.g. a .txt resume)
        # rather than rejecting the upload outright.
        text = _extract_plain_text(file_bytes)

    return text[:_MAX_TEXT_CHARS]


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return ""


def _extract_docx(file_bytes: bytes) -> str:
    try:
        document = Document(BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        # Resumes often use tables for skills/experience layout — read those too.
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
        return "\n".join(paragraphs)
    except Exception:
        return ""


def _extract_plain_text(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def detect_skills(resume_text: str) -> list[str]:
    """Match the resume text against the same known-skills vocabulary
    used for job matching (app.services.career.SKILLS), so both sides
    of a resume-vs-job comparison speak the same language."""
    text = (resume_text or "").lower()
    return sorted(skill for skill in SKILLS if skill in text)
